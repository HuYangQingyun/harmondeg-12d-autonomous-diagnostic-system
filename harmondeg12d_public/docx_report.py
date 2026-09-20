
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

OPENING=("Harmondeg 12-D Autonomous Diagnostic System v1.0 — Research Reference Result. "
"This report is generated under the Harmondeg 12-D standard diagnostic configuration and is provided for research and reference only. "
"It is not a legal, regulatory, judicial, institutional, or universally binding determination, and it does not constitute a final determination of truth or falsity.")
CLOSING=("Research Reference Only. The findings in this report reflect Harmondeg 12-D Autonomous Diagnostic System v1.0, "
"the defined diagnostic scope, operating mode, and evidence available for this analysis. Subsequent interpretation and use remain the responsibility of the user.")

def render_docx(data,out_path):
    doc=Document()
    sec=doc.sections[0]
    sec.top_margin=Inches(.7); sec.bottom_margin=Inches(.7); sec.left_margin=Inches(.75); sec.right_margin=Inches(.75)
    doc.styles["Normal"].font.name="Aptos"; doc.styles["Normal"].font.size=Pt(10)
    t=doc.add_paragraph(); t.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=t.add_run("Harmondeg 12-D Autonomous Structural Diagnostic Report"); r.bold=True; r.font.size=Pt(20)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Autonomous Diagnostic System v1.0 · Research Reference Only").italic=True
    p=doc.add_paragraph(); p.add_run("RESEARCH REFERENCE RESULT\n").bold=True; p.add_run(OPENING)

    doc.add_heading("1. Diagnostic Configuration",1)
    meta=[("Framework",data["framework_version"]),("Target",data["target"]),("Scope",data["scope"]),
          ("Operating Mode",data["operating_mode"]),("External Evidence Used","Yes" if data["external_evidence_used"] else "No"),
          ("Evidence Basis",data["evidence_basis"]),("Evidence Packet Hash",data["evidence_packet_hash"])]
    for k,v in meta:
        p=doc.add_paragraph(); p.add_run(k+": ").bold=True; p.add_run(str(v))

    doc.add_heading("2. Twelve-Dimensional Diagnosis",1)
    for k,v in data["dimensions"].items():
        doc.add_heading(f"{k} — {v.get('name','')}",2)
        p=doc.add_paragraph()
        if v["status"]=="scored":
            p.add_run("Score: ").bold=True; p.add_run(f"{v['score']:+d}")
            p.add_run("    Confidence: ").bold=True; p.add_run(str(v["confidence"]))
            for lab,key in [("Evidence","evidence"),("Structural Observation","structural_observation"),
                            ("Structural Impact","structural_impact"),("Rationale","score_rationale")]:
                q=doc.add_paragraph(); q.add_run(lab+": ").bold=True; q.add_run(v.get(key,""))
        else:
            p.add_run("Status: ").bold=True; p.add_run(v["status"])
            p.add_run("    Confidence: ").bold=True; p.add_run(str(v.get("confidence") or "—"))
            q=doc.add_paragraph(); q.add_run("Rationale: ").bold=True; q.add_run(v.get("rationale",""))

    doc.add_heading("3. Cross-Dimensional Adjudication",1)
    a=data["cross_dimensional_adjudication"]
    for lab,key in [("Source Structural Conditions","source_conditions"),("Derived Structural Effects","derived_effects"),
                    ("Associated Conditions","associated_conditions")]:
        p=doc.add_paragraph(); p.add_run(lab+": ").bold=True; p.add_run(str(a.get(key) or "None established"))
    p=doc.add_paragraph(); p.add_run("Propagation: ").bold=True
    if a.get("propagation"):
        for x in a["propagation"]:
            doc.add_paragraph(f"{x['from']} → {x['to']} · {x['strength']} · {x['basis']}",style="List Bullet")
    else: p.add_run("None established")

    doc.add_heading("4. Integrated Structural Conclusion",1)
    labels=[("theory_position","Theory Position"),("foundational_evidential_basis","Foundational / Evidential Basis"),
    ("supported_explanatory_domain","Supported Explanatory Domain"),("application_boundary","Application Boundary"),
    ("major_explanatory_strengths","Major Structural Strengths"),("unresolved_structural_conditions","Unresolved Structural Conditions"),
    ("evidential_status","Evidential Status"),("overall_structural_profile","Overall Structural Profile"),
    ("reference_use_boundary","Reference-Use Boundary")]
    for key,label in labels:
        doc.add_heading(label,2); doc.add_paragraph(data["integrated_conclusion"][key])

    doc.add_heading("5. Diagnostic Self-Audit",1)
    p=doc.add_paragraph(); p.add_run("Passed: ").bold=True; p.add_run("Yes" if data["self_audit"]["passed"] else "No")
    p=doc.add_paragraph(); p.add_run("Findings: ").bold=True; p.add_run(str(data["self_audit"]["findings"] or "No rule-level audit violations detected."))

    vals=[v["score"] for v in data["dimensions"].values() if v["status"]=="scored"]
    doc.add_heading("6. Supplementary Structural Indicator",1)
    doc.add_paragraph(f"Raw score across {len(vals)} scored dimensions: {sum(vals):+d}. Supplementary only; it does not determine the integrated diagnosis and is not an overall verdict.")

    doc.add_heading("7. Capability Boundary",1)
    doc.add_paragraph("This v1.0 contains a deterministic autonomous baseline. Where the baseline cannot establish a deep semantic distinction with sufficient reliability, it returns IE rather than manufacturing a conclusion. A compatible semantic-intelligence adapter may deepen analysis without changing the canonical Harmondeg D1–D12 rules.")

    doc.add_heading("8. Closing Research Reference Notice",1)
    doc.add_paragraph(CLOSING)
    f=doc.sections[0].footer.paragraphs[0]; f.alignment=WD_ALIGN_PARAGRAPH.CENTER
    f.add_run("Harmondeg 12-D Autonomous Diagnostic System v1.0 · Research Reference Only").font.size=Pt(8)
    doc.save(out_path)
